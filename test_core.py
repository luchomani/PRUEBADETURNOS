# -*- coding: utf-8 -*-
"""Pruebas del núcleo. Ejecutar:  python -m pytest -q test_core.py   (o  python test_core.py)"""
import json
import os
import tempfile
import threading
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook
from io import BytesIO

import zfc_core as core


@pytest.fixture(autouse=True)
def entorno(tmp_path, monkeypatch):
    for v in ("ZFC_ADMIN_PASS", "ZFC_VIG_PASS", "ZFC_ADMIN_USER", "ZFC_VIG_USER", "ZFC_SECRET_KEY"):
        monkeypatch.delenv(v, raising=False)
    core.configurar_rutas(tmp_path)
    yield tmp_path


def reg(placa="ABC123", cedula="1234567890", nombre="juan perez", tel="3001234567"):
    r, errs = core.validar_registro(placa, nombre, cedula, tel, "obs")
    assert not errs, errs
    return r


# ---------------- validación y escape ----------------
def test_validacion_ok_y_normalizacion():
    r, e = core.validar_registro(" abc-123 ", "  juan   perez ", "1.234.567", "300 123-4567", "hola\x00 mundo")
    assert not e
    assert r["placa"] == "ABC123" and r["nombre"] == "Juan Perez"
    assert r["cedula"] == "1234567" and r["telefono"] == "3001234567"
    assert "\x00" not in r["observaciones_conductor"]


@pytest.mark.parametrize("placa", ["", "AB1", "<script>", "ABC 123456789", "ÁBC123"])
def test_placa_invalida(placa):
    _, e = core.validar_registro(placa, "Juan", "1234567", "3001234", "")
    assert e


def test_cedula_telefono_invalidos():
    _, e = core.validar_registro("ABC123", "Juan", "12ab", "123", "")
    assert len(e) == 2
    _, e = core.validar_registro("ABC123", "Juan", "١٢٣٤٥٦٧", "3001234567", "")  # dígitos unicode
    assert e


def test_esc_html():
    assert core.esc('<img src=x onerror="a()">') == "&lt;img src=x onerror=&quot;a()&quot;&gt;"
    assert core.esc(None) == ""


# ---------------- datos: básico, numeración, duplicados ----------------
def test_crear_turno_numeracion_y_duplicado():
    st1, t1 = core.crear_turno(reg("AAA111"))
    st2, t2 = core.crear_turno(reg("BBB222"))
    assert (st1, t1["id_turno"]) == ("ok", "T-001")
    assert (st2, t2["id_turno"]) == ("ok", "T-002")
    st3, t3 = core.crear_turno(reg("aaa-111"))
    assert st3 == "duplicado" and t3["uid"] == t1["uid"]
    assert len(core.cargar_datos()) == 2


def test_placa_reutilizable_tras_finalizar():
    _, t = core.crear_turno(reg("AAA111"))
    core.actualizar_turno(t["uid"], {"estado": core.AUT_ING})
    assert core.confirmar_movimiento(t["uid"], "ingreso")[0] == "ok"
    assert core.confirmar_movimiento(t["uid"], "ingreso")[0] == "estado_cambio"  # ya no está autorizado
    core.actualizar_turno(t["uid"], {"estado": core.AUT_SAL})
    assert core.confirmar_movimiento(t["uid"], "salida")[0] == "ok"
    assert core.crear_turno(reg("AAA111"))[0] == "ok"


def test_buscar_activo_exige_cedula():
    core.crear_turno(reg("AAA111", cedula="1234567"))
    assert core.buscar_activo("aaa111", "1234567")
    assert core.buscar_activo("aaa111", "7654321") is None
    assert core.buscar_activo("zzz999") is None


# ---------------- concurrencia ----------------
def test_concurrencia_sin_perdida_ni_duplicados():
    n = 30
    errores = []

    def worker(i):
        try:
            core.crear_turno(reg(f"CAR{i:03d}", cedula=f"10000{i:03d}"))
        except Exception as e:  # pragma: no cover
            errores.append(e)

    hilos = [threading.Thread(target=worker, args=(i,)) for i in range(n)]
    [h.start() for h in hilos]
    [h.join() for h in hilos]
    assert not errores
    datos = core.cargar_datos()
    assert len(datos) == n
    fichos = [t["id_turno"] for t in datos]
    assert len(set(fichos)) == n, "fichos duplicados"


def test_concurrencia_vigilante_y_conductores():
    _, t = core.crear_turno(reg("AAA111"))
    core.actualizar_turno(t["uid"], {"estado": core.AUT_ING})
    hilos = [threading.Thread(target=core.crear_turno, args=(reg(f"XYZ{i:03d}", cedula=f"2000{i:04d}"),))
             for i in range(10)]
    hilos.append(threading.Thread(target=core.confirmar_movimiento, args=(t["uid"], "ingreso")))
    [h.start() for h in hilos]
    [h.join() for h in hilos]
    datos = core.cargar_datos()
    assert len(datos) == 11
    assert core.buscar_turno(datos, t["uid"])["estado"] == core.EN_ZF


def test_bloqueo_huerfano_se_rompe(entorno):
    core.LOCK_FILE.write_text("x")
    viejo = time.time() - 120
    os.utime(core.LOCK_FILE, (viejo, viejo))
    assert core.crear_turno(reg())[0] == "ok"


# ---------------- corrupción ----------------
def test_archivo_corrupto_no_se_sobrescribe_y_usa_respaldo(entorno):
    core.crear_turno(reg("AAA111"))
    core.crear_turno(reg("BBB222"))  # genera .bak con 1 registro
    core.DATA_FILE.write_text("{ esto no es json")
    datos = core.cargar_datos()
    assert len(datos) == 1  # recuperado desde .bak
    assert any(p.name.startswith("turnos_zfc.json.corrupto_") for p in entorno.iterdir())


def test_corrupto_sin_respaldo_devuelve_vacio_pero_conserva_archivo(entorno):
    core.DATA_FILE.write_text("basura")
    assert core.cargar_datos() == []
    assert any(".corrupto_" in p.name for p in entorno.iterdir())


def test_uid_legado_estable(entorno):
    core.DATA_FILE.write_text(json.dumps([{"id_turno": "T-001", "placa": "AAA111", "fecha_hora": "2026-01-01 10:00:00"}]))
    a = core.cargar_datos()[0]["uid"]
    b = core.cargar_datos()[0]["uid"]
    assert a == b


# ---------------- edición del operador ----------------
def test_editar_en_zona_franca_no_resetea_estado():
    _, t = core.crear_turno(reg())
    core.actualizar_turno(t["uid"], {"estado": core.AUT_ING})
    core.confirmar_movimiento(t["uid"], "ingreso")
    r = core.actualizar_turno(t["uid"], {"mensaje_admin": "Muelle 3"})
    assert r["estado"] == "ok" and r["despues"]["estado"] == core.EN_ZF


def test_tipo_movimiento_derivado_y_temporizador():
    _, t = core.crear_turno(reg())
    r1 = core.actualizar_turno(t["uid"], {"estado": core.AUT_ING, "mostrar_en_pantalla": True})
    assert r1["despues"]["tipo_movimiento"] == "Ingresar" and r1["despues"]["fecha_autorizacion"]
    # forzar fecha vieja y pasar a SALIR: debe reiniciarse el temporizador
    def viejo(datos):
        datos[0]["fecha_autorizacion"] = "2020-01-01 00:00:00"
        return None, True
    core.modificar_datos(viejo)
    r2 = core.actualizar_turno(t["uid"], {"estado": core.AUT_SAL})
    assert r2["despues"]["tipo_movimiento"] == "Salir"
    assert r2["despues"]["fecha_autorizacion"] > "2025-01-01"


def test_duplicados_en_edicion():
    _, a = core.crear_turno(reg("AAA111", cedula="1111111"))
    _, b = core.crear_turno(reg("BBB222", cedula="2222222"))
    assert core.actualizar_turno(b["uid"], {"placa": "AAA111"})["estado"] == "placa_duplicada"
    assert core.actualizar_turno(b["uid"], {"id_turno": a["id_turno"]})["estado"] == "ficho_duplicado"
    assert core.actualizar_turno(b["uid"], {"estado": "HACKEADO"})["estado"] == "estado_invalido"
    assert core.actualizar_turno("no-existe", {"placa": "ZZZ999"})["estado"] == "no_existe"


def test_campos_no_editables_se_ignoran():
    _, t = core.crear_turno(reg())
    r = core.actualizar_turno(t["uid"], {"uid": "otro", "fecha_hora": "1999-01-01 00:00:00", "mensaje_admin": "ok"})
    assert r["despues"]["uid"] == t["uid"] and r["despues"]["fecha_hora"] == t["fecha_hora"]


# ---------------- pantalla pública ----------------
def test_visible_en_pantalla():
    cfg = {"tiempo_en_pantalla_min": 30}
    ahora = core.ahora_str()
    base = {"mostrar_en_pantalla": True, "estado": core.AUT_ING, "fecha_autorizacion": ahora}
    assert core.visible_en_pantalla(base, cfg)
    assert not core.visible_en_pantalla({**base, "mostrar_en_pantalla": False}, cfg)
    assert not core.visible_en_pantalla({**base, "estado": core.EN_ZF}, cfg)
    assert not core.visible_en_pantalla({**base, "fecha_autorizacion": "2020-01-01 00:00:00"}, cfg)
    assert not core.visible_en_pantalla({**base, "fecha_autorizacion": None, "fecha_hora": None}, cfg)


# ---------------- eliminar / limpiar / deshacer ----------------
def test_deshacer_update_con_conflicto():
    _, t = core.crear_turno(reg())
    r = core.actualizar_turno(t["uid"], {"mensaje_admin": "A"})
    entry = {"accion": "update", "antes": r["antes"], "despues": r["despues"]}
    core.actualizar_turno(t["uid"], {"mensaje_admin": "B"})  # otro cambio posterior
    assert core.aplicar_deshacer(entry) == "conflicto"
    assert core.cargar_datos()[0]["mensaje_admin"] == "B"


def test_deshacer_update_ok_no_toca_otros():
    _, a = core.crear_turno(reg("AAA111", cedula="1111111"))
    r = core.actualizar_turno(a["uid"], {"mensaje_admin": "A"})
    _, b = core.crear_turno(reg("BBB222", cedula="2222222"))  # registro nuevo posterior
    assert core.aplicar_deshacer({"accion": "update", "antes": r["antes"], "despues": r["despues"]}) == "ok"
    datos = core.cargar_datos()
    assert len(datos) == 2 and core.buscar_turno(datos, a["uid"])["mensaje_admin"] != "A"


def test_eliminar_y_deshacer():
    _, a = core.crear_turno(reg("AAA111", cedula="1111111"))
    _, b = core.crear_turno(reg("BBB222", cedula="2222222"))
    antes = core.eliminar_turno(a["uid"])
    assert len(core.cargar_datos()) == 1
    assert core.aplicar_deshacer({"accion": "delete", "antes": antes}) == "ok"
    assert [t["id_turno"] for t in core.cargar_datos()] == ["T-001", "T-002"]
    assert core.aplicar_deshacer({"accion": "delete", "antes": antes}) == "ya_existe"


def test_limpiar_archiva_y_deshacer_conserva_nuevos(entorno):
    core.crear_turno(reg("AAA111", cedula="1111111"))
    core.crear_turno(reg("BBB222", cedula="2222222"))
    eliminado = core.limpiar_historial()
    assert len(eliminado) == 2 and core.cargar_datos() == []
    assert len(list((entorno / "archivo").glob("turnos_*.json"))) == 1
    core.crear_turno(reg("CCC333", cedula="3333333"))  # registro posterior a la limpieza
    assert core.aplicar_deshacer({"accion": "clear", "antes": eliminado}) == "ok"
    assert len(core.cargar_datos()) == 3


def test_auditoria():
    core.crear_turno(reg())
    ev = core.leer_auditoria()
    assert ev and ev[0]["accion"] == "registro"


# ---------------- contraseñas / config / tokens ----------------
def test_hash_password():
    h = core.hash_password("Secreta-123")
    assert h.startswith("pbkdf2_sha256$") and "Secreta-123" not in h
    assert core.verificar_password("Secreta-123", h)
    assert not core.verificar_password("otra", h)
    assert not core.verificar_password("x", "basura")
    assert not core.verificar_password("x", "")
    assert core.hash_password("a") != core.hash_password("a")  # sal aleatoria


def test_config_inicial_sin_texto_plano(entorno):
    cfg = core.cargar_config()
    assert cfg["cambiar_pass_admin"] is True  # contraseña por defecto -> aviso
    contenido = core.CONFIG_FILE.read_text()
    assert "ZONAFRANCA123-\"" not in contenido.replace(cfg["usuario_admin"], "")  # solo aparece como usuario
    assert "contrasena_admin" not in contenido
    assert core.verificar_credenciales(cfg, "operador", "ZONAFRANCA123-", "ZONAFRANCA123-")
    assert not core.verificar_credenciales(cfg, "operador", "ZONAFRANCA123-", "mala")
    assert not core.verificar_credenciales(cfg, "vigilante", "ZONAFRANCA123-", "ZONAFRANCA123-")
    assert core.cargar_config()["secret_key"] == cfg["secret_key"]  # estable entre cargas


def test_migracion_desde_config_antigua(entorno):
    core.CONFIG_FILE.write_text(json.dumps({
        "config_version": 2, "usuario_admin": "boss", "contrasena_admin": "MiClaveSegura9",
        "usuario_vigilante": "guarda", "contrasena_vigilante": "VIGILANTE123-",
        "bodegas": ["X", "Y"], "tiempo_en_pantalla_min": 45}))
    cfg = core.cargar_config()
    assert "contrasena_admin" not in cfg and "contrasena_admin" not in core.CONFIG_FILE.read_text()
    assert core.verificar_credenciales(cfg, "operador", "boss", "MiClaveSegura9")
    assert cfg["cambiar_pass_admin"] is False and cfg["cambiar_pass_vigilante"] is True
    assert cfg["bodegas"] == ["X", "Y"] and cfg["tiempo_en_pantalla_min"] == 45  # no se pisa lo del usuario


def test_credenciales_por_entorno(monkeypatch):
    monkeypatch.setenv("ZFC_ADMIN_PASS", "DesdeEntorno-77")
    monkeypatch.setenv("ZFC_ADMIN_USER", "admin_env")
    cfg = core.cargar_config()
    assert core.verificar_credenciales(cfg, "operador", "admin_env", "DesdeEntorno-77")
    assert cfg["cambiar_pass_admin"] is False


def test_tokens():
    s = "secreto"
    tok = core.crear_token(s, "operador")
    assert core.validar_token(s, "operador", tok)
    assert core.validar_token(s, "vigilante", tok) is None            # rol distinto
    assert core.validar_token("otro", "operador", tok) is None         # secreto distinto
    assert core.validar_token(s, "operador", "123.abc") is None        # firma falsa
    assert core.validar_token(s, "operador", "basura") is None
    assert core.validar_token(s, "operador", "") is None
    # el ataque de la versión anterior: timestamp futuro / forjado
    futuro = int(time.time()) + 10**9
    assert core.validar_token(s, "operador", f"{futuro}.{'0'*64}") is None
    assert core.validar_token(s, "operador", core.crear_token(s, "operador", futuro)) is None
    # expirado
    viejo = core.crear_token(s, "operador", int(time.time()) - core.SESION_TTL - 5)
    assert core.validar_token(s, "operador", viejo) is None


def test_limite_intentos():
    lim = core.LimiteIntentos(3, 100)
    k = ("operador", "1.2.3.4")
    for i in range(3):
        assert lim.segundos_bloqueo(k, t=1000 + i) == 0
        lim.fallo(k, t=1000 + i)
    assert lim.segundos_bloqueo(k, t=1003) > 0
    assert lim.segundos_bloqueo(("operador", "9.9.9.9"), t=1003) == 0    # otra IP no afectada
    assert lim.segundos_bloqueo(k, t=1000 + 2 + 101) == 0                # expira
    lim.fallo(k, t=2000); lim.exito(k)
    assert lim.segundos_bloqueo(k, t=2001) == 0


# ---------------- geocerca ----------------
def test_geocerca():
    cfg = {"zfc_lat": 7.9075, "zfc_lon": -72.50472, "radio_geofence_km": 3.0}
    ok, d = core.evaluar_geocerca(7.9080, -72.5050, cfg)
    assert ok and d < 0.1
    ok, d = core.evaluar_geocerca(4.7110, -74.0721, cfg)  # Bogotá
    assert not ok and 300 < d < 450


# ---------------- Excel ----------------
def test_excel_inyeccion_de_formulas():
    t = {"uid": "1", "id_turno": "T-001", "placa": "AAA111", "nombre": "=HYPERLINK(\"http://x\",\"a\")",
         "observaciones_conductor": "@SUM(A1)", "cedula": "1", "estado": core.ESPERA, "fecha_hora": "2026-01-01 10:00:00"}
    core.migrar_campos(t)
    df = core.df_exportacion([t])
    wb = load_workbook(BytesIO(core.generar_excel(df)))
    ws = wb["Reporte"]
    cabeceras = [c.value for c in ws[1]]
    celda = ws.cell(row=2, column=cabeceras.index("Conductor") + 1)
    assert celda.value.startswith("=HYPERLINK") and celda.data_type != "f"
    obs = ws.cell(row=2, column=cabeceras.index("Obs. Conductor") + 1)
    assert obs.data_type != "f"


def test_excel_vacio_y_caracteres_ilegales():
    assert load_workbook(BytesIO(core.generar_excel(core.df_exportacion([])))) is not None
    t = {"uid": "1", "id_turno": "T-001", "placa": "AAA111", "nombre": "Ana\x07", "estado": core.ESPERA}
    core.migrar_campos(t)
    core.generar_excel(core.df_exportacion([t]))  # no debe lanzar IllegalCharacterError


def test_resumen():
    core.crear_turno(reg("AAA111", cedula="1111111"))
    _, b = core.crear_turno(reg("BBB222", cedula="2222222"))
    core.actualizar_turno(b["uid"], {"estado": core.AUT_ING})
    r = core.resumen(core.cargar_datos())
    assert r == {"total": 2, "espera": 1, "autorizados": 1, "en_zf": 0, "finalizados": 0}


if __name__ == "__main__":
    raise SystemExit(pytest.main(["-q", __file__]))

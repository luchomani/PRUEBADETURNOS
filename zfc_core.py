# -*- coding: utf-8 -*-
"""
zfc_core.py — Lógica de negocio del Sistema de Turnos y Accesos (ZFC).

No depende de Streamlit: se puede probar de forma aislada.

Contiene:
  · Capa de datos con bloqueo de archivo, escritura atómica, respaldo (.bak)
    y recuperación ante corrupción.
  · Operaciones de negocio transaccionales (crear turno, confirmar ingreso /
    salida, editar, eliminar, limpiar, deshacer).
  · Configuración, contraseñas con hash (PBKDF2), tokens de sesión firmados
    (HMAC) y limitador de intentos de login.
  · Validación y saneamiento de entradas, escape de HTML.
  · Exportación a Excel (con protección contra inyección de fórmulas).
"""
from __future__ import annotations

import copy
import hashlib
import hmac
import html
import json
import math
import os
import re
import secrets
import shutil
import threading
import time
import uuid
from datetime import datetime, timedelta, timezone
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:  # QR opcional
    import qrcode
    HAS_QR = True
except ImportError:  # pragma: no cover
    HAS_QR = False

# ==========================================================================
# RUTAS
# ==========================================================================
BASE_DIR: Path
DATA_FILE: Path
CONFIG_FILE: Path
LOCK_FILE: Path
AUDIT_FILE: Path
ARCHIVE_DIR: Path


def configurar_rutas(base) -> None:
    """Define dónde viven los archivos de datos (útil para tests o despliegues)."""
    global BASE_DIR, DATA_FILE, CONFIG_FILE, LOCK_FILE, AUDIT_FILE, ARCHIVE_DIR
    BASE_DIR = Path(base)
    DATA_FILE = BASE_DIR / "turnos_zfc.json"
    CONFIG_FILE = BASE_DIR / "config_zfc.json"
    LOCK_FILE = BASE_DIR / ".zfc.lock"
    AUDIT_FILE = BASE_DIR / "auditoria_zfc.jsonl"
    ARCHIVE_DIR = BASE_DIR / "archivo"


# Por defecto junto al código; se puede mover con ZFC_DATA_DIR (p. ej. a un volumen persistente).
configurar_rutas(os.environ.get("ZFC_DATA_DIR") or Path(__file__).resolve().parent)

# ==========================================================================
# CONSTANTES
# ==========================================================================
TZ_CO = timezone(timedelta(hours=-5), "COT")  # Colombia: UTC-5 sin horario de verano
FMT_FECHA = "%Y-%m-%d %H:%M:%S"
HORA_CORTE_JORNADA = 18

ESPERA = "En Espera"
AUT_ING = "AUTORIZADO PARA INGRESAR"
AUT_SAL = "AUTORIZADO PARA SALIR"
EN_ZF = "EN ZONA FRANCA"
FINAL = "FINALIZADO"
ESTADOS = [ESPERA, AUT_ING, AUT_SAL, EN_ZF, FINAL]

SESION_TTL = 7200  # 2 horas
PBKDF2_ITER = 200_000
MAX_HISTORIAL_UNDO = 20

PASSWORDS_POR_DEFECTO = {"ZONAFRANCA123-", "VIGILANTE123-"}

DEFAULT_CONFIG = {
    "config_version": 3,
    "tiempo_en_pantalla_min": 30,
    "usuario_admin": "ZONAFRANCA123-",
    "usuario_vigilante": "VIGILANTE123-",
    "hash_admin": "",
    "hash_vigilante": "",
    "cambiar_pass_admin": False,
    "cambiar_pass_vigilante": False,
    "secret_key": "",
    "url_base": "",
    "geocerca_activa": False,
    "radio_geofence_km": 3.0,
    "zfc_lat": 7.90750,
    "zfc_lon": -72.50472,
    "bodegas": [
        "OTRA BODEGA",
        "ALFRANCA CUCUTA S.A.S.",
        "ALMACODER SAS",
        "ALMANORTE S.A.S",
        "BERZATEX S.A.S.",
        "COSMOGREEN SAS",
        "DEL BEN SAS",
        "GLEX S.A.S.",
        "GRUPO EMPRESARIAL SUAPEL S.A.S.",
        "HSQ GROUP INTERNATIONAL SAS",
        "INDULOGISTICS SAS",
        "INTERNATIONAL FRONTIER S.A.S.",
        "INVERSIONES CARGAMAX HYJ S.A.S.",
        "LOGISTICA ADUANERA J.H. LTDA",
        "LOGISTICA COEX SAS",
        "M.B. INDUSTRIAS Y SERVICIOS SAS",
        "MANUFEX GROUP S.A.S.",
        "SETRACOMEX S.A.S.",
        "TRENDS LOVERS S.A.S.",
        "VEHIMAQ ZONA FRANCA SAS",
    ],
}

# rol -> (clave usuario, clave hash, bandera "cambiar", legado texto plano,
#         variable entorno usuario, variable entorno clave, clave por defecto)
_ROLES = {
    "operador": ("usuario_admin", "hash_admin", "cambiar_pass_admin",
                 "contrasena_admin", "ZFC_ADMIN_USER", "ZFC_ADMIN_PASS", "ZONAFRANCA123-"),
    "vigilante": ("usuario_vigilante", "hash_vigilante", "cambiar_pass_vigilante",
                  "contrasena_vigilante", "ZFC_VIG_USER", "ZFC_VIG_PASS", "VIGILANTE123-"),
}

# ==========================================================================
# TIEMPO
# ==========================================================================


def ahora() -> datetime:
    return datetime.now(TZ_CO)


def ahora_str() -> str:
    return ahora().strftime(FMT_FECHA)


def ultimo_corte(ref: datetime | None = None) -> datetime:
    """Inicio de la jornada actual (corte diario a las 18:00 hora Colombia)."""
    ref = ref or ahora()
    corte = ref.replace(hour=HORA_CORTE_JORNADA, minute=0, second=0, microsecond=0)
    return corte if ref.hour >= HORA_CORTE_JORNADA else corte - timedelta(days=1)


def ultimo_corte_str(ref: datetime | None = None) -> str:
    return ultimo_corte(ref).strftime(FMT_FECHA)


def parse_fecha(s) -> datetime | None:
    if not s:
        return None
    try:
        return datetime.strptime(s, FMT_FECHA).replace(tzinfo=TZ_CO)
    except (ValueError, TypeError):
        return None


def fecha_legible(s) -> str:
    if not s:
        return "—"
    try:
        return datetime.strptime(s, FMT_FECHA).strftime("%d/%m/%Y %H:%M:%S")
    except (ValueError, TypeError):
        return str(s)


# ==========================================================================
# TEXTO / VALIDACIÓN / ESCAPE
# ==========================================================================
_RE_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
RE_PLACA = re.compile(r"^[A-Z0-9]{5,8}$")
RE_FICHO = re.compile(r"^[A-Z0-9][A-Z0-9\-]{0,11}$")
_RE_TURNO = re.compile(r"^T-(\d+)$")


def esc(v) -> str:
    """Escapa HTML. Úsalo SIEMPRE antes de insertar datos de usuario en HTML."""
    return html.escape("" if v is None else str(v), quote=True)


def limpiar_texto(s, max_len: int = 100, multilinea: bool = False) -> str:
    s = "" if s is None else str(s)
    s = _RE_CTRL.sub("", s)
    if multilinea:
        s = "\n".join(" ".join(linea.split()) for linea in s.splitlines()).strip()
    else:
        s = " ".join(s.split())
    return s[:max_len]


def norm_placa(p) -> str:
    return re.sub(r"[\s\-]", "", (p or "").upper())


def validar_campo(campo: str, valor) -> tuple[str, str | None]:
    """Devuelve (valor_limpio, mensaje_error|None)."""
    if campo == "placa":
        v = norm_placa(limpiar_texto(valor, 20))
        if not v:
            return v, "La placa del vehículo es obligatoria."
        if not RE_PLACA.match(v):
            return v, "La placa debe tener entre 5 y 8 caracteres (solo letras y números)."
        return v, None
    if campo == "nombre":
        v = limpiar_texto(valor, 80).title()
        if not v:
            return v, "El nombre del conductor es obligatorio."
        if len(v) < 3:
            return v, "El nombre debe tener al menos 3 caracteres."
        return v, None
    if campo == "cedula":
        v = re.sub(r"[\s.\-]", "", limpiar_texto(valor, 20))
        if not v:
            return v, "El número de cédula es obligatorio."
        if not re.fullmatch(r"[0-9]+", v):
            return v, "La cédula debe contener solo números."
        if not 5 <= len(v) <= 15:
            return v, "La cédula debe tener entre 5 y 15 dígitos."
        return v, None
    if campo == "telefono":
        v = re.sub(r"[\s\-()+]", "", limpiar_texto(valor, 25))
        if not v:
            return v, "El teléfono de contacto es obligatorio."
        if not re.fullmatch(r"[0-9]+", v):
            return v, "El teléfono debe contener solo números."
        if not 7 <= len(v) <= 15:
            return v, "El teléfono debe tener entre 7 y 15 dígitos."
        return v, None
    if campo == "ficho":
        v = limpiar_texto(valor, 12).upper()
        if not v:
            return v, "El número de Ficho no puede estar vacío."
        if not RE_FICHO.match(v):
            return v, "El Ficho solo admite letras, números y guion (máx. 12 caracteres)."
        return v, None
    raise ValueError(f"Campo desconocido: {campo}")


def validar_registro(placa, nombre, cedula, telefono, observaciones="") -> tuple[dict, list[str]]:
    """Valida el formulario del conductor. Devuelve (datos_limpios, errores)."""
    reg, errores = {}, []
    for campo, raw in (("placa", placa), ("nombre", nombre), ("cedula", cedula), ("telefono", telefono)):
        limpio, err = validar_campo(campo, raw)
        reg[campo] = limpio
        if err:
            errores.append(err)
    reg["observaciones_conductor"] = limpiar_texto(observaciones, 300, multilinea=True)
    return reg, errores


# ==========================================================================
# ESTADOS
# ==========================================================================


def estado(t: dict) -> str:
    return (t.get("estado") or "").strip().upper()


def es(t: dict, *estados: str) -> bool:
    return estado(t) in {e.upper() for e in estados}


def es_final(t: dict) -> bool:
    return es(t, FINAL)


def badge_html(estado_raw) -> str:
    e = (estado_raw or "").strip().upper()
    if e == AUT_ING:
        return '<span class="badge badge-ingresar">Autorizado para Ingresar</span>'
    if e == AUT_SAL:
        return '<span class="badge badge-salir">Autorizado para Salir</span>'
    if e == EN_ZF:
        return '<span class="badge badge-zf">En Zona Franca</span>'
    if e == FINAL:
        return '<span class="badge badge-finalizado">Finalizado</span>'
    return '<span class="badge badge-espera">En Espera</span>'


def visible_en_pantalla(t: dict, cfg: dict) -> bool:
    """¿Debe mostrarse este turno en la pantalla pública ahora mismo?"""
    if not t.get("mostrar_en_pantalla"):
        return False
    if not es(t, AUT_ING, AUT_SAL):
        return False
    ref = parse_fecha(t.get("fecha_autorizacion")) or parse_fecha(t.get("fecha_hora"))
    if ref is None:
        return False
    limite = timedelta(minutes=cfg.get("tiempo_en_pantalla_min", 30))
    return (ahora() - ref) <= limite


def resumen(turnos: list[dict]) -> dict:
    return {
        "total": len(turnos),
        "espera": sum(es(t, ESPERA) for t in turnos),
        "autorizados": sum(es(t, AUT_ING, AUT_SAL) for t in turnos),
        "en_zf": sum(es(t, EN_ZF) for t in turnos),
        "finalizados": sum(es_final(t) for t in turnos),
    }


# ==========================================================================
# CAPA DE DATOS: bloqueo + escritura atómica + respaldo
# ==========================================================================
class ErrorBloqueo(RuntimeError):
    pass


@contextmanager
def _bloqueo(timeout: float = 10.0, obsoleto: float = 30.0):
    """Bloqueo entre procesos/hilos con archivo .lock (portable)."""
    BASE_DIR.mkdir(parents=True, exist_ok=True)
    inicio = time.time()
    fd = None
    while fd is None:
        try:
            fd = os.open(str(LOCK_FILE), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            try:  # bloqueo huérfano (proceso caído): se rompe
                if time.time() - os.path.getmtime(LOCK_FILE) > obsoleto:
                    os.remove(LOCK_FILE)
                    continue
            except OSError:
                pass
            if time.time() - inicio > timeout:
                raise ErrorBloqueo("No se pudo obtener el bloqueo de datos.")
            time.sleep(0.03)
    try:
        yield
    finally:
        try:
            os.close(fd)
        finally:
            try:
                os.remove(LOCK_FILE)
            except OSError:
                pass


def _bak(path: Path) -> Path:
    return path.with_name(path.name + ".bak")


def _escritura_atomica(path: Path, objeto) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f"{path.name}.{os.getpid()}.{threading.get_ident()}.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(objeto, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    if path.exists():
        try:
            shutil.copy2(path, _bak(path))
        except OSError:
            pass
    os.replace(tmp, path)


def _leer_lista(path: Path) -> list:
    with open(path, "r", encoding="utf-8") as f:
        datos = json.load(f)
    if not isinstance(datos, list):
        raise ValueError("El archivo de turnos no contiene una lista.")
    return datos


def _aislar_corrupto(path: Path) -> None:
    """Aparta (no borra) un archivo corrupto para poder recuperarlo a mano."""
    try:
        destino = path.with_name(f"{path.name}.corrupto_{datetime.now().strftime('%Y%m%d_%H%M%S')}")
        os.replace(path, destino)
    except OSError:
        pass


def _defaults_turno() -> dict:
    return {
        "id_turno": "", "placa": "", "nombre": "", "cedula": "", "telefono": "",
        "bodega": "", "tipo_movimiento": "Ingresar", "estado": ESPERA,
        "fecha_hora": "", "mensaje_admin": "", "mostrar_en_pantalla": False,
        "fecha_autorizacion": None, "fecha_ingreso_zf": None, "fecha_salida_zf": None,
        "checklist_ingreso": {}, "checklist_salida": {},
        "observaciones_vigilante": "", "observaciones_conductor": "",
    }


def migrar_campos(item: dict) -> None:
    for k, v in _defaults_turno().items():
        item.setdefault(k, v)
    if not item.get("uid"):
        # uid determinista: estable entre lecturas aunque aún no se haya guardado
        base = f"zfc|{item['id_turno']}|{item['placa']}|{item['fecha_hora']}"
        item["uid"] = str(uuid.uuid5(uuid.NAMESPACE_URL, base))


def cargar_datos() -> list[dict]:
    """Lee los turnos. Si el archivo está corrupto lo aparta y usa el respaldo.
    Nunca sobrescribe un archivo corrupto con una lista vacía."""
    datos = None
    for ruta in (DATA_FILE, _bak(DATA_FILE)):
        if not ruta.exists():
            continue
        try:
            datos = _leer_lista(ruta)
            break
        except (OSError, ValueError):
            if ruta == DATA_FILE:
                _aislar_corrupto(ruta)
    if datos is None:
        return []
    items = [d for d in datos if isinstance(d, dict)]
    for it in items:
        migrar_campos(it)
    return items


def modificar_datos(fn):
    """Transacción: bajo bloqueo carga los datos, ejecuta fn(datos) y guarda
    solo si fn devuelve cambiado=True.  fn debe devolver (resultado, cambiado)."""
    with _bloqueo():
        datos = cargar_datos()
        resultado, cambiado = fn(datos)
        if cambiado:
            _escritura_atomica(DATA_FILE, datos)
        return resultado


def _clave_orden(t: dict) -> tuple:
    return (str(t.get("fecha_hora", "")), str(t.get("id_turno", "")))


def buscar_turno(turnos: list[dict], uid) -> dict | None:
    return next((t for t in turnos if t.get("uid") == uid), None)


def buscar_activo(placa: str, cedula: str | None = None) -> dict | None:
    """Turno no finalizado para una placa (y cédula si se indica)."""
    p = norm_placa(placa)
    if not p:
        return None
    for t in reversed(cargar_datos()):
        if norm_placa(t.get("placa")) == p and not es_final(t):
            if cedula is None or re.sub(r"[\s.\-]", "", str(cedula)) == str(t.get("cedula", "")):
                return t
    return None


# ==========================================================================
# AUDITORÍA
# ==========================================================================
def registrar_auditoria(rol: str, accion: str, turno: str = "", detalle: str = "") -> None:
    try:
        BASE_DIR.mkdir(parents=True, exist_ok=True)
        linea = {"fecha": ahora_str(), "rol": rol, "accion": accion,
                 "turno": turno, "detalle": str(detalle)[:400]}
        with open(AUDIT_FILE, "a", encoding="utf-8") as f:
            f.write(json.dumps(linea, ensure_ascii=False) + "\n")
    except OSError:
        pass


def leer_auditoria(n: int = 200) -> list[dict]:
    try:
        with open(AUDIT_FILE, "r", encoding="utf-8") as f:
            lineas = f.readlines()[-n:]
    except OSError:
        return []
    out = []
    for ln in reversed(lineas):
        try:
            out.append(json.loads(ln))
        except ValueError:
            continue
    return out


# ==========================================================================
# OPERACIONES DE NEGOCIO (todas transaccionales)
# ==========================================================================
def crear_turno(reg: dict) -> tuple[str, dict]:
    """Registro de un conductor. Devuelve ("ok"|"duplicado", turno)."""
    def _op(datos):
        placa = reg["placa"]
        activo = next((t for t in datos
                       if norm_placa(t.get("placa")) == placa and not es_final(t)), None)
        if activo:
            return ("duplicado", copy.deepcopy(activo)), False
        corte = ultimo_corte_str()
        nums = []
        for t in datos:
            if str(t.get("fecha_hora", "")) >= corte:
                m = _RE_TURNO.match(str(t.get("id_turno", "")))
                if m:
                    nums.append(int(m.group(1)))
        nuevo = {
            "uid": str(uuid.uuid4()),
            "id_turno": f"T-{(max(nums) + 1 if nums else 1):03d}",
            "placa": placa,
            "nombre": reg["nombre"],
            "cedula": reg["cedula"],
            "telefono": reg["telefono"],
            "bodega": reg.get("bodega", ""),
            "tipo_movimiento": "Ingresar",
            "estado": ESPERA,
            "fecha_hora": ahora_str(),
            "mensaje_admin": "En espera de autorización",
            "mostrar_en_pantalla": False,
            "fecha_autorizacion": None,
            "fecha_ingreso_zf": None,
            "fecha_salida_zf": None,
            "checklist_ingreso": {},
            "checklist_salida": {},
            "observaciones_vigilante": "",
            "observaciones_conductor": reg.get("observaciones_conductor", ""),
            "acepta_datos_fecha": ahora_str(),
        }
        datos.append(nuevo)
        return ("ok", copy.deepcopy(nuevo)), True

    res = modificar_datos(_op)
    if res[0] == "ok":
        registrar_auditoria("conductor", "registro", res[1]["id_turno"], res[1]["placa"])
    return res


def confirmar_movimiento(uid: str, tipo: str, observacion: str = "") -> tuple[str, dict | None]:
    """Vigilante confirma ingreso o salida. tipo: 'ingreso' | 'salida'.
    Comprueba bajo bloqueo que el estado siga siendo el esperado."""
    if tipo not in ("ingreso", "salida"):
        raise ValueError("tipo debe ser 'ingreso' o 'salida'")
    esperado = AUT_ING if tipo == "ingreso" else AUT_SAL
    obs = limpiar_texto(observacion, 300, multilinea=True)

    def _op(datos):
        t = buscar_turno(datos, uid)
        if t is None:
            return ("no_existe", None), False
        if not es(t, esperado):
            return ("estado_cambio", copy.deepcopy(t)), False
        hora = ahora_str()
        if tipo == "ingreso":
            t["fecha_ingreso_zf"] = hora
            t["estado"] = EN_ZF
        else:
            t["fecha_salida_zf"] = hora
            t["estado"] = FINAL
        t["mostrar_en_pantalla"] = False
        if obs:
            t["observaciones_vigilante"] = obs
        return ("ok", copy.deepcopy(t)), True

    res = modificar_datos(_op)
    if res[0] == "ok":
        registrar_auditoria("vigilante", f"confirmar_{tipo}", res[1]["id_turno"], res[1]["placa"])
    return res


_CAMPOS_EDITABLES = ("id_turno", "placa", "nombre", "cedula", "telefono", "bodega",
                     "estado", "mensaje_admin", "mostrar_en_pantalla")


def actualizar_turno(uid: str, cambios: dict) -> dict:
    """Edición del operador. Devuelve {"estado": status, "antes":..., "despues":...}.
    status: ok | no_existe | placa_duplicada | ficho_duplicado | estado_invalido"""
    cambios = {k: v for k, v in cambios.items() if k in _CAMPOS_EDITABLES}

    def _op(datos):
        t = buscar_turno(datos, uid)
        if t is None:
            return {"estado": "no_existe"}, False

        nuevo_estado = cambios.get("estado", t.get("estado"))
        if "estado" in cambios and nuevo_estado not in ESTADOS:
            return {"estado": "estado_invalido"}, False

        placa = norm_placa(cambios.get("placa", t.get("placa")))
        if nuevo_estado.upper() != FINAL and any(
            norm_placa(o.get("placa")) == placa and o.get("uid") != uid and not es_final(o)
            for o in datos
        ):
            return {"estado": "placa_duplicada"}, False

        ficho = cambios.get("id_turno", t.get("id_turno"))
        corte = ultimo_corte_str()
        if any(o.get("id_turno") == ficho and str(o.get("fecha_hora", "")) >= corte
               and o.get("uid") != uid for o in datos):
            return {"estado": "ficho_duplicado"}, False

        antes = copy.deepcopy(t)
        antes_estado = estado(t)
        antes_mostrar = bool(t.get("mostrar_en_pantalla"))

        t.update(cambios)
        nuevo_up = str(t.get("estado", "")).upper()
        # tipo_movimiento se deriva del estado (evita contradicciones)
        if nuevo_up == AUT_ING:
            t["tipo_movimiento"] = "Ingresar"
        elif nuevo_up == AUT_SAL:
            t["tipo_movimiento"] = "Salir"
        # el temporizador de pantalla se reinicia al autorizar, al cambiar de
        # autorización o al activar "mostrar en pantalla"
        if nuevo_up in (AUT_ING, AUT_SAL) and (
            nuevo_up != antes_estado
            or (t.get("mostrar_en_pantalla") and not antes_mostrar)
            or not t.get("fecha_autorizacion")
        ):
            t["fecha_autorizacion"] = ahora_str()

        return {"estado": "ok", "antes": antes, "despues": copy.deepcopy(t)}, True

    res = modificar_datos(_op)
    if res["estado"] == "ok":
        dif = {k: [res["antes"].get(k), res["despues"].get(k)]
               for k in _CAMPOS_EDITABLES if res["antes"].get(k) != res["despues"].get(k)}
        registrar_auditoria("operador", "editar", res["despues"]["id_turno"],
                            json.dumps(dif, ensure_ascii=False))
    return res


def eliminar_turno(uid: str) -> dict | None:
    def _op(datos):
        for i, t in enumerate(datos):
            if t.get("uid") == uid:
                return datos.pop(i), True
        return None, False

    antes = modificar_datos(_op)
    if antes:
        registrar_auditoria("operador", "eliminar", antes.get("id_turno", ""), antes.get("placa", ""))
    return antes


def limpiar_historial() -> list[dict]:
    """Archiva TODO el historial en /archivo y lo vacía. Devuelve lo eliminado."""
    def _op(datos):
        if not datos:
            return [], False
        ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
        destino = ARCHIVE_DIR / f"turnos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        _escritura_atomica(destino, datos)
        eliminado = copy.deepcopy(datos)
        datos.clear()
        return eliminado, True

    eliminado = modificar_datos(_op)
    if eliminado:
        registrar_auditoria("operador", "limpiar_historial", "", f"{len(eliminado)} registros archivados")
    return eliminado


def aplicar_deshacer(entry: dict) -> str:
    """Deshace UNA acción concreta sin tocar el resto de los datos.
    Devuelve: ok | conflicto | no_existe | ya_existe | invalido"""
    accion = entry.get("accion")

    def _op(datos):
        if accion == "update":
            antes, despues = entry["antes"], entry["despues"]
            for i, t in enumerate(datos):
                if t.get("uid") == antes["uid"]:
                    if t != despues:  # alguien más lo modificó después
                        return "conflicto", False
                    datos[i] = copy.deepcopy(antes)
                    return "ok", True
            return "no_existe", False
        if accion == "delete":
            antes = entry["antes"]
            if buscar_turno(datos, antes["uid"]):
                return "ya_existe", False
            datos.append(copy.deepcopy(antes))
            datos.sort(key=_clave_orden)
            return "ok", True
        if accion == "clear":
            existentes = {t.get("uid") for t in datos}
            faltan = [copy.deepcopy(t) for t in entry["antes"] if t.get("uid") not in existentes]
            if not faltan:
                return "ya_existe", False
            datos.extend(faltan)
            datos.sort(key=_clave_orden)
            return "ok", True
        return "invalido", False

    res = modificar_datos(_op)
    if res == "ok":
        registrar_auditoria("operador", f"deshacer_{accion}", "", "")
    return res


# ==========================================================================
# CONFIGURACIÓN, CONTRASEÑAS, TOKENS, LÍMITE DE INTENTOS
# ==========================================================================
def hash_password(password: str, iteraciones: int = PBKDF2_ITER) -> str:
    salt = secrets.token_bytes(16)
    dk = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, iteraciones)
    return f"pbkdf2_sha256${iteraciones}${salt.hex()}${dk.hex()}"


def verificar_password(password: str, almacenado: str) -> bool:
    try:
        alg, it, salt_hex, hash_hex = (almacenado or "").split("$")
        if alg != "pbkdf2_sha256":
            return False
        dk = hashlib.pbkdf2_hmac("sha256", (password or "").encode("utf-8"),
                                 bytes.fromhex(salt_hex), int(it))
        return hmac.compare_digest(dk.hex(), hash_hex)
    except (ValueError, TypeError):
        return False


def validar_password_nueva(p: str) -> str | None:
    if len(p) < 8:
        return "La contraseña debe tener al menos 8 caracteres."
    if p in PASSWORDS_POR_DEFECTO:
        return "No uses la contraseña por defecto."
    return None


def _leer_config_archivo() -> dict:
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            cfg = json.load(f)
        return cfg if isinstance(cfg, dict) else {}
    except (OSError, ValueError):
        return {}


def _construir_config() -> tuple[dict, bool]:
    """Mezcla archivo + defaults, migra contraseñas en texto plano a hash y
    genera la clave secreta. Devuelve (cfg, hubo_cambios)."""
    guardada = _leer_config_archivo()
    cfg = copy.deepcopy(DEFAULT_CONFIG)
    cfg.update(guardada)
    cambio = False

    for rol, (ukey, hkey, flag, legado, env_u, env_p, defecto) in _ROLES.items():
        plano = guardada.get(legado)
        if plano is not None:  # migración desde la versión anterior
            cfg[hkey] = hash_password(str(plano))
            cfg[flag] = str(plano) in PASSWORDS_POR_DEFECTO
            cfg.pop(legado, None)
            cambio = True
        elif not cfg.get(hkey):
            pwd = os.environ.get(env_p) or defecto
            if os.environ.get(env_u):
                cfg[ukey] = os.environ[env_u]
            cfg[hkey] = hash_password(pwd)
            cfg[flag] = pwd in PASSWORDS_POR_DEFECTO
            cambio = True

    if not cfg.get("secret_key"):
        cfg["secret_key"] = secrets.token_hex(32)
        cambio = True
    if cfg.get("config_version") != DEFAULT_CONFIG["config_version"]:
        cfg["config_version"] = DEFAULT_CONFIG["config_version"]
        cambio = True
    if not isinstance(cfg.get("bodegas"), list):
        cfg["bodegas"] = list(DEFAULT_CONFIG["bodegas"])
        cambio = True
    return cfg, cambio


def _guardar_config_sin_bloqueo(cfg: dict) -> None:
    _escritura_atomica(CONFIG_FILE, cfg)


def guardar_config(cfg: dict) -> None:
    with _bloqueo():
        _guardar_config_sin_bloqueo(cfg)


def cargar_config() -> dict:
    cfg, cambio = _construir_config()
    if cambio:
        with _bloqueo():
            cfg, cambio = _construir_config()  # otro proceso pudo adelantarse
            if cambio:
                _guardar_config_sin_bloqueo(cfg)
    return cfg


def rotar_secreto(cfg: dict) -> None:
    """Invalida todos los tokens de sesión emitidos (p. ej. al cambiar contraseñas)."""
    cfg["secret_key"] = secrets.token_hex(32)


def secret_de(cfg: dict) -> str:
    return os.environ.get("ZFC_SECRET_KEY") or cfg["secret_key"]


def verificar_credenciales(cfg: dict, rol: str, usuario: str, password: str) -> bool:
    ukey, hkey = _ROLES[rol][0], _ROLES[rol][1]
    u_ok = hmac.compare_digest((usuario or "").encode("utf-8"), str(cfg.get(ukey, "")).encode("utf-8"))
    p_ok = verificar_password(password or "", cfg.get(hkey, ""))  # siempre se calcula
    return u_ok and p_ok


def _firmar(secret: str, rol: str, ts: int) -> str:
    return hmac.new(secret.encode("utf-8"), f"{rol}|{ts}".encode("utf-8"), hashlib.sha256).hexdigest()


def crear_token(secret: str, rol: str, ts: int | None = None) -> str:
    ts = int(time.time()) if ts is None else int(ts)
    return f"{ts}.{_firmar(secret, rol, ts)}"


def validar_token(secret: str, rol: str, token: str, ahora_ts: float | None = None) -> int | None:
    """Devuelve el timestamp de emisión si el token es válido y vigente; si no, None."""
    try:
        ts_s, firma = str(token).split(".", 1)
        ts = int(ts_s)
    except (ValueError, TypeError):
        return None
    ahora_ts = time.time() if ahora_ts is None else ahora_ts
    if not (0 <= ahora_ts - ts < SESION_TTL):
        return None
    if not hmac.compare_digest(firma, _firmar(secret, rol, ts)):
        return None
    return ts


class LimiteIntentos:
    """Bloquea una clave (p. ej. rol+IP) tras N fallos dentro de una ventana."""

    def __init__(self, max_intentos: int = 5, ventana_s: int = 300):
        self.max = max_intentos
        self.ventana = ventana_s
        self._fallos: dict = {}
        self._lock = threading.Lock()

    def _recientes(self, clave, t):
        lista = [x for x in self._fallos.get(clave, []) if t - x < self.ventana]
        if lista:
            self._fallos[clave] = lista
        else:
            self._fallos.pop(clave, None)
        return lista

    def segundos_bloqueo(self, clave, t: float | None = None) -> int:
        t = time.time() if t is None else t
        with self._lock:
            lista = self._recientes(clave, t)
            if len(lista) >= self.max:
                return max(1, int(self.ventana - (t - lista[-1])) + 1)
            return 0

    def fallo(self, clave, t: float | None = None) -> None:
        t = time.time() if t is None else t
        with self._lock:
            self._recientes(clave, t)
            self._fallos.setdefault(clave, []).append(t)

    def exito(self, clave) -> None:
        with self._lock:
            self._fallos.pop(clave, None)


# ==========================================================================
# GEOCERCA
# ==========================================================================
def distancia_km(lat1, lon1, lat2, lon2) -> float:
    r = 6371.0088
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = p2 - p1, math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def evaluar_geocerca(lat, lon, cfg: dict) -> tuple[bool, float]:
    d = distancia_km(float(lat), float(lon), float(cfg["zfc_lat"]), float(cfg["zfc_lon"]))
    return d <= float(cfg["radio_geofence_km"]), d


# ==========================================================================
# EXPORTACIÓN
# ==========================================================================
_COLUMNAS_EXPORT = [
    ("id_turno", "Turno"), ("placa", "Placa"), ("nombre", "Conductor"),
    ("cedula", "Cédula"), ("telefono", "Teléfono"), ("fecha_hora", "Fecha Registro"),
    ("bodega", "Bodega"), ("tipo_movimiento", "Movimiento"), ("estado", "Estado"),
    ("mensaje_admin", "Mensaje"), ("fecha_autorizacion", "Autorizado Operador"),
    ("fecha_ingreso_zf", "Ingreso Vigilante"), ("fecha_salida_zf", "Salida Vigilante"),
    ("observaciones_conductor", "Obs. Conductor"), ("observaciones_vigilante", "Obs. Vigilante"),
]
_CAMPOS_FECHA = {"fecha_hora", "fecha_autorizacion", "fecha_ingreso_zf", "fecha_salida_zf"}


def df_exportacion(turnos: list[dict]) -> pd.DataFrame:
    filas = []
    for t in turnos:
        filas.append({
            etiqueta: (fecha_legible(t.get(k)) if k in _CAMPOS_FECHA else t.get(k, ""))
            for k, etiqueta in _COLUMNAS_EXPORT
        })
    return pd.DataFrame(filas, columns=[e for _, e in _COLUMNAS_EXPORT])


def _sanear_df(df: pd.DataFrame) -> pd.DataFrame:
    def _f(v):
        return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v
    aplicar = getattr(df, "map", None) or df.applymap
    return aplicar(_f)


def generar_excel(df: pd.DataFrame) -> bytes:
    df = _sanear_df(df)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Reporte")
        ws = writer.sheets["Reporte"]

        header_fill = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        lado = Side(style="thin", color="D9D9D9")
        borde = Border(left=lado, right=lado, top=lado, bottom=lado)

        for idx, nombre in enumerate(df.columns, start=1):
            celda = ws.cell(row=1, column=idx)
            celda.fill, celda.font, celda.border = header_fill, header_font, borde
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            largo = max([len(str(nombre))] + [len(str(v)) for v in df.iloc[:, idx - 1]])
            ws.column_dimensions[get_column_letter(idx)].width = min(largo + 3, 40)

        for fila in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(df.columns)):
            for celda in fila:
                celda.border = borde
                celda.alignment = Alignment(vertical="top", wrap_text=True)

        # Anti-inyección de fórmulas: todo texto que empiece como fórmula se fuerza a texto
        for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(df.columns)):
            for celda in fila:
                if isinstance(celda.value, str) and celda.value[:1] in ("=", "+", "-", "@"):
                    celda.data_type = "s"
        ws.freeze_panes = "A2"
    return buffer.getvalue()


def generar_qr_png(url: str) -> bytes | None:
    if not HAS_QR or not url:
        return None
    try:
        qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_H,
                           box_size=15, border=4)
        qr.add_data(url)
        qr.make(fit=True)
        img = qr.make_image(fill_color="#000000", back_color="#FFFFFF")
        buf = BytesIO()
        img.save(buf, format="PNG")
        return buf.getvalue()
    except Exception:
        return None
